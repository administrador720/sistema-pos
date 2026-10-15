"""Exportación de informes a Excel y PDF, y recibo de venta en PDF (80 mm)."""
from __future__ import annotations

import io
from datetime import datetime

from .. import config


def _fmt(valor, tipo: str) -> str:
    if valor is None or valor == "":
        return ""
    if tipo == "dinero":
        return "$ " + f"{int(round(valor)):,}".replace(",", ".")
    if tipo == "entero":
        return f"{int(valor):,}".replace(",", ".")
    if tipo == "porcentaje":
        return f"{valor:.1f} %".replace(".", ",")
    if tipo == "decimal":
        return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if tipo in ("fecha", "fechahora"):
        s = str(valor).replace("T", " ")
        if tipo == "fecha":
            return s[:10]
        return s[:16]
    return str(valor)


def a_xlsx(informe: dict, negocio: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = informe["titulo"][:31].replace("/", "-").replace(":", "-")
    ws.append([negocio]) if negocio else None
    ws.append([informe["titulo"]])
    periodo = f"{informe.get('desde', '')} a {informe.get('hasta', '')}" if informe.get("desde") else ""
    ws.append([periodo, f"Generado: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    cols = informe["columnas"]
    inicio = ws.max_row + 1
    ws.append([c["titulo"] for c in cols])
    for c in ws[inicio]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    formatos = {"dinero": '"$" #,##0', "entero": "#,##0", "porcentaje": '0.0"%"', "decimal": "#,##0.00"}
    for fila in informe["filas"]:
        ws.append([fila.get(c["key"]) for c in cols])
    for j, c in enumerate(cols, start=1):
        fmt = formatos.get(c["tipo"])
        for r in range(inicio + 1, ws.max_row + 1):
            celda = ws.cell(row=r, column=j)
            if fmt:
                celda.number_format = fmt
    if informe.get("totales"):
        ws.append(["Total" if i == 0 else informe["totales"].get(c["key"]) for i, c in enumerate(cols)])
        for celda in ws[ws.max_row]:
            celda.font = Font(bold=True)
        for j, c in enumerate(cols, start=1):
            if formatos.get(c["tipo"]):
                ws.cell(row=ws.max_row, column=j).number_format = formatos[c["tipo"]]
    if informe.get("nota"):
        ws.append([])
        ws.append([informe["nota"]])
    for j, c in enumerate(cols, start=1):
        ancho = max([len(c["titulo"])] + [len(_fmt(f.get(c["key"]), c["tipo"])) for f in informe["filas"][:200]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 10), 48)
    ws.freeze_panes = ws.cell(row=inicio + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def a_pdf(informe: dict, negocio: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    cols = informe["columnas"]
    horizontal = len(cols) > 6
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if horizontal else A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=informe["titulo"])
    est = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=est["Heading1"], fontSize=16, textColor=colors.HexColor("#0F766E"), spaceAfter=2)
    peq = ParagraphStyle("p", parent=est["Normal"], fontSize=8, textColor=colors.HexColor("#6B7280"))
    celda = ParagraphStyle("c", parent=est["Normal"], fontSize=7.5, leading=9)
    historia = []
    if negocio:
        historia.append(Paragraph(negocio, peq))
    historia.append(Paragraph(informe["titulo"], titulo))
    periodo = f"Periodo: {informe.get('desde', '')} a {informe.get('hasta', '')} · " if informe.get("desde") else ""
    historia.append(Paragraph(f"{periodo}Generado: {datetime.now():%Y-%m-%d %H:%M}", peq))
    historia.append(Spacer(1, 4 * mm))
    datos = [[Paragraph(f"<b>{c['titulo']}</b>", ParagraphStyle("h", parent=celda, textColor=colors.white)) for c in cols]]
    for f in informe["filas"][:3000]:
        datos.append([Paragraph(_fmt(f.get(c["key"]), c["tipo"]).replace("&", "&amp;").replace("<", "&lt;"), celda) for c in cols])
    if informe.get("totales"):
        tot = informe["totales"]
        datos.append([Paragraph("<b>" + ("Total" if i == 0 else _fmt(tot.get(c["key"]), c["tipo"])) + "</b>", celda)
                      for i, c in enumerate(cols)])
    tabla = Table(datos, repeatRows=1)
    estilo = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
              ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
              ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    if informe.get("totales"):
        estilo.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E6F2F1")))
    tabla.setStyle(TableStyle(estilo))
    historia.append(tabla)
    if len(informe["filas"]) > 3000:
        historia.append(Paragraph("Se muestran las primeras 3.000 filas; exporte a Excel para ver todas.", peq))
    if informe.get("nota"):
        historia += [Spacer(1, 3 * mm), Paragraph(informe["nota"], peq)]
    doc.build(historia)
    return buf.getvalue()


def recibo_pdf(venta: dict, negocio: dict) -> bytes:
    """Tiquete de venta en ancho de 80 mm (para imprimir cuando haya impresora)."""
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    ancho = 80 * mm
    lineas_items = len(venta["items"]) * 2 + len(venta["pagos"]) + 18
    alto = max(120 * mm, lineas_items * 4.2 * mm)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(ancho, alto))
    y = alto - 8 * mm

    def centro(txt, tam=9, negrita=False):
        nonlocal y
        c.setFont("Helvetica-Bold" if negrita else "Helvetica", tam)
        c.drawCentredString(ancho / 2, y, txt)
        y -= (tam + 3)

    def par(izq, der, tam=8, negrita=False):
        nonlocal y
        c.setFont("Helvetica-Bold" if negrita else "Helvetica", tam)
        c.drawString(4 * mm, y, izq)
        c.drawRightString(ancho - 4 * mm, y, der)
        y -= (tam + 3)

    def linea():
        nonlocal y
        c.setDash(1, 2)
        c.line(4 * mm, y + 2, ancho - 4 * mm, y + 2)
        c.setDash()
        y -= 5

    def din(v):
        return "$ " + f"{int(v):,}".replace(",", ".")

    centro(negocio.get("negocio_nombre", ""), 11, True)
    for k in ("negocio_nit", "negocio_direccion", "negocio_telefono"):
        if negocio.get(k):
            centro(("NIT " if k == "negocio_nit" else "") + negocio[k], 7)
    linea()
    par(f"Tiquete {venta['numero']}", venta["fecha"][:16].replace("T", " "), 8, True)
    par("Cajero", venta["cajero"] or "")
    if venta.get("estado") == "anulada":
        centro("*** ANULADA ***", 11, True)
    linea()
    for it in venta["items"]:
        c.setFont("Helvetica", 8)
        c.drawString(4 * mm, y, it["descripcion"][:34])
        y -= 10
        par(f"  {it['cantidad']} x {din(it['precio_unitario'])}", din(it["total"] + 0))
    linea()
    par("Subtotal", din(venta["subtotal"]))
    if venta["descuento_total"]:
        par("Descuento", "-" + din(venta["descuento_total"]))
    par("Impuestos incluidos", din(venta["impuesto_total"]))
    par("TOTAL", din(venta["total"]), 11, True)
    linea()
    for p in venta["pagos"]:
        par(p["medio"], din(p["monto"]))
        if p["tipo"] == "efectivo" and p["devuelta"]:
            par("  Recibido", din(p["recibido"]), 7)
            par("  Devuelta", din(p["devuelta"]), 7, True)
    y -= 4
    centro("¡Gracias por su compra!", 8)
    c.showPage()
    c.save()
    return buf.getvalue()


def guardar_en_exportes(nombre: str, contenido: bytes) -> str:
    """Escribe el archivo en la carpeta de exportes (para la ventana de escritorio)."""
    destino = config.dir_exportes() / nombre
    destino.write_bytes(contenido)
    return str(destino)
