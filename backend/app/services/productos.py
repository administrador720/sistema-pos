"""Productos: alta/edición con códigos de barras múltiples, historial de precios, cambio masivo e importación."""
from __future__ import annotations

import csv
import io
import re
import unicodedata

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from .. import models as m
from ..core import ErrorNegocio, ahora, auditar, redondear
from . import inventario as inv


def serializar(p: m.Producto, *, completo: bool = True) -> dict:
    d = {
        "id": p.id, "codigo_interno": p.codigo_interno, "nombre": p.nombre, "precio_venta": p.precio_venta,
        "stock_actual": p.stock_actual, "activo": p.activo, "categoria_id": p.categoria_id,
        "categoria": p.categoria.nombre if p.categoria else None, "ubicacion": p.ubicacion,
        "impuesto_id": p.impuesto_id,
        "impuesto": p.impuesto.nombre if p.impuesto else None,
        "tarifa_pct": p.impuesto.tarifa_pct if p.impuesto else 0,
        "codigos": [{"codigo": c.codigo, "unidades": c.unidades_por_codigo} for c in p.codigos],
    }
    if completo:
        d.update({
            "marca": p.marca, "proveedor_id": p.proveedor_id, "proveedor": p.proveedor.nombre if p.proveedor else None,
            "costo_promedio": round(p.costo_promedio, 2), "stock_minimo": p.stock_minimo, "stock_maximo": p.stock_maximo,
            "margen_pct": round((1 - (p.costo_promedio / (p.precio_venta / (1 + (p.impuesto.tarifa_pct if p.impuesto else 0) / 100))))
                                * 100, 1) if p.precio_venta else None,
            "creado_en": p.creado_en.isoformat(), "actualizado_en": p.actualizado_en.isoformat(),
        })
    return d


def _opciones():
    return (selectinload(m.Producto.codigos), selectinload(m.Producto.categoria),
            selectinload(m.Producto.impuesto), selectinload(m.Producto.proveedor))


def por_codigo(db: Session, codigo: str) -> tuple[m.Producto, int] | None:
    cb = db.get(m.CodigoBarra, codigo.strip())
    if not cb:
        # el código interno también sirve para vender
        p = db.scalars(select(m.Producto).where(m.Producto.codigo_interno == codigo.strip())
                       .options(*_opciones())).first()
        return (p, 1) if p else None
    p = db.scalars(select(m.Producto).where(m.Producto.id == cb.producto_id).options(*_opciones())).one()
    return p, cb.unidades_por_codigo


def _normalizar(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s.lower()) if unicodedata.category(c) != "Mn")


def buscar(db: Session, q: str = "", *, categoria_id: int | None = None, proveedor_id: int | None = None,
           estado: str = "activos", bajo_minimo: bool = False, limite: int = 100, offset: int = 0) -> dict:
    qry = select(m.Producto).options(*_opciones())
    cnt = select(func.count()).select_from(m.Producto)
    conds = []
    if q.strip():
        for palabra in q.split():
            like = f"%{palabra}%"
            conds.append(or_(m.Producto.nombre.ilike(like), m.Producto.marca.ilike(like),
                             m.Producto.codigo_interno.ilike(like),
                             m.Producto.id.in_(select(m.CodigoBarra.producto_id).where(m.CodigoBarra.codigo.like(f"{palabra}%")))))
    if categoria_id:
        conds.append(m.Producto.categoria_id == categoria_id)
    if proveedor_id:
        conds.append(m.Producto.proveedor_id == proveedor_id)
    if estado == "activos":
        conds.append(m.Producto.activo.is_(True))
    elif estado == "inactivos":
        conds.append(m.Producto.activo.is_(False))
    if bajo_minimo:
        conds.append(m.Producto.stock_actual <= m.Producto.stock_minimo)
    for c in conds:
        qry, cnt = qry.where(c), cnt.where(c)
    total = db.scalar(cnt) or 0
    filas = db.scalars(qry.order_by(m.Producto.nombre).limit(limite).offset(offset)).all()
    return {"total": total, "productos": [serializar(p) for p in filas]}


def _validar(datos: dict, parcial: bool = False):
    if not parcial or "nombre" in datos:
        if not (datos.get("nombre") or "").strip():
            raise ErrorNegocio("El nombre del producto es obligatorio.")
    for campo in ("precio_venta", "stock_minimo", "stock_maximo"):
        if campo in datos and datos[campo] is not None and int(datos[campo]) < 0:
            raise ErrorNegocio(f"{campo.replace('_', ' ').capitalize()} no puede ser negativo.")
    if "costo_promedio" in datos and datos["costo_promedio"] is not None and float(datos["costo_promedio"]) < 0:
        raise ErrorNegocio("El costo no puede ser negativo.")


def _asignar_codigos(db: Session, p: m.Producto, codigos: list[dict]):
    nuevos: dict[str, int] = {}
    for c in codigos:
        cod = str(c.get("codigo", "")).strip()
        if not cod:
            continue
        uni = int(c.get("unidades") or 1)
        if uni < 1:
            raise ErrorNegocio("Las unidades por código deben ser al menos 1.")
        nuevos[cod] = uni
    for cod in nuevos:
        otro = db.get(m.CodigoBarra, cod)
        if otro and otro.producto_id != p.id:
            raise ErrorNegocio(f"El código {cod} ya pertenece a «{db.get(m.Producto, otro.producto_id).nombre}».", 409)
    actuales = {c.codigo: c for c in p.codigos}
    for cod, c in list(actuales.items()):
        if cod not in nuevos:
            p.codigos.remove(c)
    db.flush()
    for cod, uni in nuevos.items():
        if cod in actuales:
            actuales[cod].unidades_por_codigo = uni
        else:
            p.codigos.append(m.CodigoBarra(codigo=cod, unidades_por_codigo=uni))


def crear(db: Session, usuario: m.Usuario, datos: dict) -> m.Producto:
    _validar(datos)
    ahora_ = ahora()
    p = m.Producto(
        codigo_interno=(datos.get("codigo_interno") or "").strip() or None, nombre=datos["nombre"].strip(),
        categoria_id=datos.get("categoria_id"), marca=(datos.get("marca") or "").strip() or None,
        proveedor_id=datos.get("proveedor_id"), impuesto_id=datos.get("impuesto_id"),
        costo_promedio=float(datos.get("costo_promedio") or 0), precio_venta=int(datos.get("precio_venta") or 0),
        stock_minimo=int(datos.get("stock_minimo") or 0), stock_maximo=int(datos.get("stock_maximo") or 0),
        ubicacion=(datos.get("ubicacion") or "").strip() or None, activo=bool(datos.get("activo", True)),
        stock_actual=0, creado_en=ahora_, actualizado_en=ahora_)
    if p.codigo_interno and db.scalar(select(m.Producto.id).where(m.Producto.codigo_interno == p.codigo_interno)):
        raise ErrorNegocio("Ya existe un producto con ese código interno.", 409)
    db.add(p)
    db.flush()
    _asignar_codigos(db, p, datos.get("codigos") or [])
    inicial = int(datos.get("stock_inicial") or 0)
    if inicial:
        inv.saldo_inicial(db, usuario.id, p, inicial, p.costo_promedio)
    auditar(db, usuario.id, "PRODUCTO_CREADO", "producto", p.id, despues={"nombre": p.nombre, "precio": p.precio_venta})
    db.flush()
    return p


def cambiar_precio(db: Session, p: m.Producto, nuevo: int, usuario_id: int | None):
    if nuevo < 0:
        raise ErrorNegocio("El precio no puede ser negativo.")
    if nuevo == p.precio_venta:
        return
    db.add(m.HistorialPrecio(producto_id=p.id, precio_anterior=p.precio_venta, precio_nuevo=nuevo,
                             usuario_id=usuario_id, fecha=ahora()))
    auditar(db, usuario_id, "CAMBIO_PRECIO", "producto", p.id, antes={"precio": p.precio_venta}, despues={"precio": nuevo})
    p.precio_venta = nuevo
    p.actualizado_en = ahora()


def actualizar(db: Session, usuario: m.Usuario, producto_id: int, datos: dict) -> m.Producto:
    p = db.scalars(select(m.Producto).where(m.Producto.id == producto_id).options(*_opciones())).first()
    if not p:
        raise ErrorNegocio("Producto no encontrado.", 404)
    _validar(datos, parcial=True)
    if "precio_venta" in datos and datos["precio_venta"] is not None:
        cambiar_precio(db, p, int(datos["precio_venta"]), usuario.id)
    for campo in ("nombre", "marca", "ubicacion", "codigo_interno"):
        if campo in datos:
            v = (datos[campo] or "").strip() or None
            if campo == "codigo_interno" and v and v != p.codigo_interno and db.scalar(
                    select(m.Producto.id).where(m.Producto.codigo_interno == v)):
                raise ErrorNegocio("Ya existe un producto con ese código interno.", 409)
            setattr(p, campo, v if campo != "nombre" else (v or p.nombre))
    for campo in ("categoria_id", "proveedor_id", "impuesto_id", "stock_minimo", "stock_maximo"):
        if campo in datos:
            setattr(p, campo, datos[campo] if datos[campo] is None or campo.endswith("_id") else int(datos[campo]))
    if "activo" in datos:
        p.activo = bool(datos["activo"])
    if "costo_promedio" in datos and datos["costo_promedio"] is not None:
        if float(datos["costo_promedio"]) != p.costo_promedio:
            auditar(db, usuario.id, "CAMBIO_COSTO", "producto", p.id, antes={"costo": p.costo_promedio},
                    despues={"costo": float(datos["costo_promedio"])})
        p.costo_promedio = float(datos["costo_promedio"])
    if "codigos" in datos:
        _asignar_codigos(db, p, datos["codigos"] or [])
    p.actualizado_en = ahora()
    db.flush()
    return p


def cambio_masivo_precios(db: Session, usuario: m.Usuario, *, porcentaje: float, categoria_id: int | None = None,
                          proveedor_id: int | None = None, redondeo: int = 50, simular: bool = True) -> dict:
    """Sube/baja precios un % (negativo baja) y redondea al múltiplo `redondeo` (p. ej. 50 pesos)."""
    if not -90 <= porcentaje <= 500:
        raise ErrorNegocio("Porcentaje fuera de rango.")
    qry = select(m.Producto).where(m.Producto.activo.is_(True), m.Producto.precio_venta > 0)
    if categoria_id:
        qry = qry.where(m.Producto.categoria_id == categoria_id)
    if proveedor_id:
        qry = qry.where(m.Producto.proveedor_id == proveedor_id)
    cambios = []
    for p in db.scalars(qry):
        nuevo = p.precio_venta * (1 + porcentaje / 100)
        nuevo = redondear(nuevo / redondeo) * redondeo if redondeo and redondeo > 1 else redondear(nuevo)
        if nuevo != p.precio_venta:
            cambios.append((p, nuevo))
    if not simular:
        for p, nuevo in cambios:
            cambiar_precio(db, p, nuevo, usuario.id)
    return {"afectados": len(cambios), "aplicado": not simular,
            "muestra": [{"id": p.id, "nombre": p.nombre, "antes": p.precio_venta, "despues": n} for p, n in cambios[:15]]}


def historial_precios(db: Session, producto_id: int) -> list[dict]:
    return [{"fecha": h.fecha.isoformat(), "anterior": h.precio_anterior, "nuevo": h.precio_nuevo, "usuario": u}
            for h, u in db.execute(select(m.HistorialPrecio, m.Usuario.nombre).outerjoin(
                m.Usuario, m.Usuario.id == m.HistorialPrecio.usuario_id).where(
                m.HistorialPrecio.producto_id == producto_id).order_by(m.HistorialPrecio.id.desc()))]


# ───────────────────────────── Importación desde Excel / CSV ─────────────────────────────
COLUMNAS = ["codigo_barras", "nombre", "categoria", "marca", "proveedor", "impuesto", "costo", "precio",
            "stock_inicial", "stock_minimo", "ubicacion"]
_ALIAS = {"codigo": "codigo_barras", "codigo de barras": "codigo_barras", "barras": "codigo_barras",
          "ean": "codigo_barras", "producto": "nombre", "descripcion": "nombre", "costo_unitario": "costo",
          "precio_venta": "precio", "stock": "stock_inicial", "existencia": "stock_inicial", "minimo": "stock_minimo"}


def plantilla_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
    ws.append(["7701234567890", "Leche entera 1 L", "Lácteos", "Alquería", "Distribuidora Lácteos", "Excluido", 3100, 4200, 24, 12, "Pasillo 3"])
    ws.append(["7709876543210", "Aceite vegetal 1 L", "Despensa", "Premier", "Grasas del Valle", "IVA 19%", 8900, 11500, 18, 6, "Pasillo 1"])
    ws.append(["", "Una línea por producto. Código y ubicación opcionales. Un código con varias unidades se agrega luego en el producto.", "", "", "", "", "", "", "", "", ""])
    for i, w in enumerate([18, 34, 18, 16, 24, 14, 10, 10, 14, 14, 14], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _leer_filas(nombre_archivo: str, contenido: bytes) -> list[dict]:
    nombre = nombre_archivo.lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
    elif nombre.endswith((".csv", ".txt")):
        texto = contenido.decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(texto[:2048], delimiters=",;\t") if texto.strip() else csv.excel
        filas = list(csv.reader(io.StringIO(texto), dialecto))
    else:
        raise ErrorNegocio("Formato no soportado. Use .xlsx o .csv.")
    if not filas:
        raise ErrorNegocio("El archivo está vacío.")
    enc = []
    for h in filas[0]:
        h = _normalizar(str(h or "")).strip().replace("  ", " ")
        enc.append(_ALIAS.get(h, h.replace(" ", "_")))
    if "nombre" not in enc:
        raise ErrorNegocio("No se encontró la columna «nombre». Descargue la plantilla para ver el formato.")
    return [{enc[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(fila) if i < len(enc)}
            for fila in filas[1:] if any(v not in (None, "") for v in fila)]


def _num(valor: str, campo: str, entero=True):
    if valor in ("", None):
        return 0
    try:
        x = float(str(valor).replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")
                  if re.search(r"\d\.\d{3}(,|$)", str(valor)) else str(valor).replace(",", "."))
    except ValueError:
        raise ValueError(f"{campo} no es un número: «{valor}»")
    if x < 0:
        raise ValueError(f"{campo} no puede ser negativo")
    return int(round(x)) if entero else x


def importar(db: Session, usuario: m.Usuario, nombre_archivo: str, contenido: bytes, simular: bool = True) -> dict:
    filas = _leer_filas(nombre_archivo, contenido)
    impuestos = {_normalizar(i.nombre): i for i in db.scalars(select(m.Impuesto))}
    creados = actualizados = 0
    errores = []
    cache_cat: dict[str, m.Categoria] = {_normalizar(c.nombre): c for c in db.scalars(select(m.Categoria))}
    cache_prov: dict[str, m.Proveedor] = {_normalizar(p.nombre): p for p in db.scalars(select(m.Proveedor))}
    vistos: set[str] = set()
    for n, f in enumerate(filas, start=2):
        try:
            nombre = f.get("nombre", "").strip()
            if not nombre:
                raise ValueError("falta el nombre")
            codigo = re.sub(r"\.0$", "", f.get("codigo_barras", "").strip())
            if codigo and codigo in vistos:
                raise ValueError(f"el código {codigo} está repetido en el archivo")
            if codigo:
                vistos.add(codigo)
            costo = _num(f.get("costo", ""), "costo", entero=False)
            precio = _num(f.get("precio", ""), "precio")
            stock = int(_num(f.get("stock_inicial", ""), "stock_inicial"))
            minimo = int(_num(f.get("stock_minimo", ""), "stock_minimo"))
            imp = None
            if f.get("impuesto"):
                imp = impuestos.get(_normalizar(f["impuesto"]))
                if not imp:
                    raise ValueError(f"impuesto «{f['impuesto']}» no existe (use: {', '.join(i.nombre for i in impuestos.values())})")
            existente = db.get(m.CodigoBarra, codigo) if codigo else None
            if existente:
                actualizados += 1
            else:
                creados += 1
            if simular:
                continue
            cat = None
            if f.get("categoria"):
                k = _normalizar(f["categoria"])
                cat = cache_cat.get(k) or cache_cat.setdefault(k, _nuevo(db, m.Categoria(nombre=f["categoria"].strip())))
            prov = None
            if f.get("proveedor"):
                k = _normalizar(f["proveedor"])
                prov = cache_prov.get(k) or cache_prov.setdefault(k, _nuevo(db, m.Proveedor(nombre=f["proveedor"].strip())))
            datos = {"nombre": nombre, "categoria_id": cat.id if cat else None, "marca": f.get("marca"),
                     "proveedor_id": prov.id if prov else None, "impuesto_id": imp.id if imp else None,
                     "costo_promedio": costo, "precio_venta": precio, "stock_minimo": minimo,
                     "ubicacion": f.get("ubicacion"), "codigos": [{"codigo": codigo, "unidades": 1}] if codigo else []}
            if existente:
                datos.pop("codigos")
                actualizar(db, usuario, existente.producto_id, datos)
            else:
                datos["stock_inicial"] = stock
                crear(db, usuario, datos)
        except ValueError as e:
            errores.append({"fila": n, "nombre": f.get("nombre", ""), "error": str(e)})
        except ErrorNegocio as e:
            errores.append({"fila": n, "nombre": f.get("nombre", ""), "error": e.detalle})
    if errores and not simular:
        db.rollback()
        raise ErrorNegocio(f"No se importó nada: {len(errores)} fila(s) con error. Corríjalas y reintente.", 422,
                           "IMPORTACION_CON_ERRORES")
    if not simular:
        auditar(db, usuario.id, "IMPORTACION_PRODUCTOS", despues={"creados": creados, "actualizados": actualizados})
    return {"filas": len(filas), "creados": creados, "actualizados": actualizados, "errores": errores,
            "aplicado": not simular}


def _nuevo(db: Session, obj):
    db.add(obj)
    db.flush()
    return obj
